from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path('Attendance_Shift_OT_QA_SOP.docx')
BLUE = '2E74B5'
DARK_BLUE = '1F4D78'
LIGHT_BLUE = 'E8EEF5'
LIGHT_GRAY = 'F2F4F7'
INK = '0B2545'
MUTED = '555555'
WHITE = 'FFFFFF'
WIDTH = 9360


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn('w:shd'))
    if shading is None:
        shading = OxmlElement('w:shd')
        properties.append(shading)
    shading.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in('w:tcMar')
    if margins is None:
        margins = OxmlElement('w:tcMar')
        properties.append(margins)
    for name, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = margins.find(qn(f'w:{name}'))
        if node is None:
            node = OxmlElement(f'w:{name}')
            margins.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_cell_width(cell, width):
    cell.width = Inches(width / 1440)
    properties = cell._tc.get_or_add_tcPr()
    width_node = properties.find(qn('w:tcW'))
    if width_node is None:
        width_node = OxmlElement('w:tcW')
        properties.append(width_node)
    width_node.set(qn('w:w'), str(width))
    width_node.set(qn('w:type'), 'dxa')


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    properties = table._tbl.tblPr
    width_node = properties.find(qn('w:tblW'))
    if width_node is None:
        width_node = OxmlElement('w:tblW')
        properties.append(width_node)
    width_node.set(qn('w:w'), str(sum(widths)))
    width_node.set(qn('w:type'), 'dxa')
    indent = properties.find(qn('w:tblInd'))
    if indent is None:
        indent = OxmlElement('w:tblInd')
        properties.append(indent)
    indent.set(qn('w:w'), '120')
    indent.set(qn('w:type'), 'dxa')
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement('w:gridCol')
        column.set(qn('w:w'), str(width))
        grid.append(column)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_borders(table, color='B7C9DF', size='6'):
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        properties.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        node = borders.find(qn(f'w:{edge}'))
        if node is None:
            node = OxmlElement(f'w:{edge}')
            borders.append(node)
        node.set(qn('w:val'), 'single')
        node.set(qn('w:sz'), size)
        node.set(qn('w:space'), '0')
        node.set(qn('w:color'), color)


def set_run_font(run, name='Calibri', size=11, color=None, bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:ascii'), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:hAnsi'), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph(paragraph, before=0, after=6, line=1.25, alignment=None):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if alignment is not None:
        paragraph.alignment = alignment


def add_text(paragraph, text, **kwargs):
    run = paragraph.add_run(text)
    set_run_font(run, **kwargs)
    return run


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f'Heading {level}')
    add_text(paragraph, text, size={1: 16, 2: 13, 3: 12}[level],
             color={1: BLUE, 2: BLUE, 3: DARK_BLUE}[level], bold=True)
    return paragraph


def add_body(doc, text, bold_prefix=None):
    paragraph = doc.add_paragraph(style='Normal')
    if bold_prefix and text.startswith(bold_prefix):
        add_text(paragraph, bold_prefix, bold=True)
        add_text(paragraph, text[len(bold_prefix):])
    else:
        add_text(paragraph, text)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style='List Bullet')
    add_text(paragraph, text)
    return paragraph


def add_number(doc, text):
    paragraph = doc.add_paragraph(style='List Number')
    add_text(paragraph, text)
    return paragraph


def add_callout(doc, label, text, fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [WIDTH])
    set_borders(table, color='A9C2DD', size='8')
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    paragraph = cell.paragraphs[0]
    set_paragraph(paragraph, after=0, line=1.15)
    add_text(paragraph, f'{label}: ', color=DARK_BLUE, bold=True)
    add_text(paragraph, text, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, headers, rows, widths, header_fill=DARK_BLUE):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_borders(table)
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, header_fill)
        paragraph = cell.paragraphs[0]
        set_paragraph(paragraph, after=0, line=1.0)
        add_text(paragraph, header, color=WHITE, size=10, bold=True)
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            if index % 2:
                set_cell_shading(cell, LIGHT_GRAY)
            paragraph = cell.paragraphs[0]
            set_paragraph(paragraph, after=0, line=1.0)
            add_text(paragraph, str(value), size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_test_case(doc, case_id, title, steps, expected, evidence):
    add_heading(doc, f'{case_id}: {title}', 2)
    add_body(doc, 'Steps')
    for step in steps:
        add_number(doc, step)
    add_body(doc, 'Expected result')
    for item in expected:
        add_bullet(doc, item)
    add_body(doc, f'Evidence: {evidence}')


def configure_styles(doc):
    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Calibri'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for level, size, color, before, after in [
        (1, 16, BLUE, 18, 10),
        (2, 13, BLUE, 14, 7),
        (3, 12, DARK_BLUE, 10, 5),
    ]:
        style = styles[f'Heading {level}']
        style.font.name = 'Calibri'
        style._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
        style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for style_name in ('List Bullet', 'List Number'):
        style = styles[style_name]
        style.font.name = 'Calibri'
        style._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
        style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def main():
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header.paragraphs[0]
    set_paragraph(header, after=0, line=1.0, alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    add_text(header, 'HR ATTENDANCE QA  |  CONTROLLED TEST DOCUMENT', size=9, color=MUTED, bold=True)
    footer = section.footer.paragraphs[0]
    set_paragraph(footer, before=0, after=0, line=1.0, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(footer, 'Attendance, Shift and Overtime QA SOP', size=9, color=MUTED)

    title = doc.add_paragraph()
    set_paragraph(title, before=12, after=3, line=1.0, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(title, 'Attendance, Shift and Overtime QA SOP', size=24, color=DARK_BLUE, bold=True)
    subtitle = doc.add_paragraph()
    set_paragraph(subtitle, after=16, line=1.0, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(subtitle, 'Standard operating procedure for functional testing in Odoo', size=11, color=MUTED, italic=True)

    add_table(doc, ['Document control', 'Value'], [
        ('Version', '1.0'),
        ('Owner', 'QA / HR Operations'),
        ('Scope', 'Attendance, shifts, leave, regularization and overtime'),
        ('Timezone', 'Asia/Kolkata'),
        ('Test mode', 'Manual attendance and biometric synchronization'),
    ], [1800, 7560], header_fill=BLUE)

    add_callout(doc, 'Purpose', 'Confirm that shift setup, attendance punches, late check-in, early checkout, overtime and leave rules produce the expected results. Overtime is checked in Payroll > Overtime Register and is not calculated in payslips.')

    add_heading(doc, '1. Preconditions', 1)
    for item in [
        'Use a dedicated QA employee and a test company.',
        'Set the QA user timezone to Asia/Kolkata.',
        'Confirm the employee is active and has the correct attendance category.',
        'Confirm the employee device user ID matches the biometric device ID.',
        'Confirm an active contract and resource calendar are available.',
        'Use a draft database or test company so production attendance is not affected.',
    ]:
        add_bullet(doc, item)

    add_heading(doc, '2. Test configuration', 1)
    add_table(doc, ['Configuration', 'QA value'], [
        ('Shift start', '09:00'),
        ('Shift end', '18:00'),
        ('Break duration', '01:00'),
        ('Check-in grace', '10 minutes'),
        ('Check-out grace', '10 minutes'),
        ('Overtime allowed', 'Yes for OT tests'),
        ('Weekly off', 'Sunday'),
    ], [3000, 6360])
    add_callout(doc, 'Expected work time', '09:00 to 18:00 with a one-hour break equals 8 payable working hours.')

    add_heading(doc, '3. Shift precedence', 1)
    add_body(doc, 'The system should select the effective shift in this order:')
    add_table(doc, ['Priority', 'Source', 'Purpose'], [
        ('1', 'Shift Roster', 'Actual daily shift plan for a specific employee and date'),
        ('2', 'Shift Assignment', 'Temporary employee, department or category assignment'),
        ('3', 'Employee Default Shift', 'Fallback shift when no roster or assignment exists'),
        ('4', 'Contract Calendar', 'Final fallback when no shift template is configured'),
    ], [900, 2700, 5760])

    add_heading(doc, '4. Functional test cases', 1)
    add_test_case(doc, 'TC-001', 'Default shift attendance',
                  ['Set the General Shift as the employee Default Shift.', 'Create attendance with check-in 09:00 and check-out 18:00.', 'Open the attendance record.'],
                  ['The employee uses the default shift.', 'Worked hours are approximately 8.00 after the one-hour break.', 'Late check-in and overtime are 0.00.'],
                  'Attendance record screenshot and values.')
    add_test_case(doc, 'TC-002', 'Shift assignment override',
                  ['Set the employee default shift to General Shift.', 'Create an assignment for 14:00 to 22:00 on the test date.', 'Generate attendance using check-in 14:16 and check-out 22:00.'],
                  ['The assignment shift is used instead of the default shift.', 'Late check-in is 6 minutes after the 10-minute grace.', 'Overtime is 0.00.'],
                  'Assignment and attendance screenshots.')
    add_test_case(doc, 'TC-003', 'Roster daily shift',
                  ['Create a roster for the test date.', 'Generate roster lines and confirm the roster.', 'Open the employee roster line.', 'Create attendance according to the roster time.'],
                  ['A daily roster line exists for the employee.', 'The roster shift is used for attendance calculations.', 'Roster takes priority over assignment and default shift.'],
                  'Roster line and attendance screenshots.')
    add_test_case(doc, 'TC-004', 'Late check-in',
                  ['Use a 09:00 shift with 10 minutes grace.', 'Test check-in at 08:55, 09:00, 09:10, 09:11 and 09:16.', 'Open each attendance record.'],
                  ['08:55, 09:00 and 09:10 are not late.', '09:11 shows 1 minute late.', '09:16 shows 6 minutes late.'],
                  'Late minutes from each record.')
    add_test_case(doc, 'TC-005', 'Early checkout',
                  ['Use an 18:00 shift end with 10 minutes grace.', 'Test checkout at 17:55, 17:49 and 18:00.', 'Open the attendance details.'],
                  ['Checkout within grace is not marked early.', 'Checkout at 17:49 calculates early leave.', 'Checkout at or after 18:00 has no early leave.'],
                  'Early leave values and screenshot.')
    add_test_case(doc, 'TC-006', 'Automatic overtime',
                  ['Enable Overtime Allowed on the shift.', 'Check in at 09:00 and check out at 20:00.', 'Synchronize or save the attendance.', 'Open Payroll > Overtime Register.'],
                  ['Scheduled work is 8.00 hours.', 'Automatic overtime is 2.00 hours.', 'The attendance date and employee are correct in Overtime Register.'],
                  'Attendance and Overtime Register screenshots.')
    add_test_case(doc, 'TC-007', 'Overtime disabled',
                  ['Disable Overtime Allowed on the shift.', 'Check in at 09:00 and check out at 20:00.', 'Open Overtime Register.'],
                  ['Overtime is 0.00 when overtime is disabled.', 'Worked hours remain calculated normally.'],
                  'Shift setting and attendance screenshot.')
    add_test_case(doc, 'TC-008', 'Manual overtime',
                  ['Open Payroll > Overtime Requests.', 'Create a request for 2.00 hours.', 'Submit and approve the request as HR Manager.', 'Open Payroll > Overtime Register.'],
                  ['Request status is Approved.', 'Approved Hours is 2.00.', 'Manual OT is visible and linked to the attendance/request.'],
                  'Approved request and register screenshots.')
    add_test_case(doc, 'TC-009', 'Biometric synchronization',
                  ['Punch check-in on the biometric device.', 'Punch check-out on the biometric device.', 'Run the device synchronization/import.', 'Search the employee and date in Attendances.'],
                  ['Both punches are created in Odoo.', 'The date and time are correct in Asia/Kolkata.', 'Worked hours, late minutes and OT are calculated after synchronization.'],
                  'Device log and Odoo attendance screenshots.')
    add_test_case(doc, 'TC-010', 'Missed checkout regularization',
                  ['Create or synchronize only a check-in.', 'Open Attendance Regularization.', 'Select Missed Check-out and the attendance date.', 'Enter a valid checkout and approve the request.'],
                  ['Existing check-in is shown.', 'Missing checkout is clearly identified.', 'Approval updates the attendance record.', 'Invalid checkout earlier than check-in is rejected.'],
                  'Regularization form and updated attendance.')
    add_test_case(doc, 'TC-011', 'Leave and public holiday',
                  ['Apply Casual, Sick, Earned, Paid and Unpaid leave on separate test dates.', 'Configure one public holiday and one weekly off.', 'Generate the roster and review attendance.'],
                  ['Leave codes remain correct: CL, SL, EL, PL and UNP.', 'Paid leave does not create LOP.', 'Unpaid leave creates LOP.', 'Public holiday and weekly off follow company policy.'],
                  'Leave records, roster and attendance screenshots.')
    add_test_case(doc, 'TC-012', 'Refresh and regression',
                  ['Create a draft payslip for the test period.', 'Change attendance or leave.', 'Click Refresh Data and then Compute Sheet.', 'Review the payslip and Overtime Register.'],
                  ['Latest attendance and leave values are loaded without deleting the payslip.', 'OT remains checked separately in Overtime Register.', 'Completed payslips are not changed automatically.'],
                  'Before/after screenshots and result values.')

    add_heading(doc, '5. Overtime Register review', 1)
    add_body(doc, 'Open Payroll > Overtime Register and verify that the combined register contains attendance-based OT and manual OT linked to attendance.')
    add_table(doc, ['Field', 'Verification'], [
        ('Employee', 'Correct employee is displayed'),
        ('Date', 'Attendance date is correct'),
        ('Check-in / Check-out', 'Punches match the device or manual entry'),
        ('Worked Hours', 'Matches shift time minus break'),
        ('Overtime Hours', 'Automatic attendance OT is visible'),
        ('Validated OT', 'Approved/validated OT is visible when used'),
        ('Manual OT', 'Approved manual OT hours are visible'),
        ('Manual OT Request', 'Linked request can be opened'),
    ], [2400, 6960])

    add_heading(doc, '6. QA evidence and result', 1)
    add_table(doc, ['Field', 'Value'], [
        ('Test Case ID', ''),
        ('Employee', ''),
        ('Date', ''),
        ('Shift / roster', ''),
        ('Check-in / check-out', ''),
        ('Worked hours', ''),
        ('Late minutes', ''),
        ('Early leave minutes', ''),
        ('Automatic OT', ''),
        ('Manual OT', ''),
        ('Expected result', ''),
        ('Actual result', ''),
        ('Status', 'PASS / FAIL'),
        ('Screenshot or evidence', ''),
    ], [3000, 6360])

    add_heading(doc, '7. Defect report format', 1)
    add_table(doc, ['Field', 'Details'], [
        ('Issue title', ''),
        ('Module', 'Attendance / Shift / Roster / Regularization / Overtime'),
        ('Test Case ID', ''),
        ('Steps to reproduce', ''),
        ('Expected result', ''),
        ('Actual result', ''),
        ('Error message', ''),
        ('Screenshot', ''),
        ('Severity', 'Low / Medium / High / Critical'),
    ], [2400, 6960])

    add_callout(doc, 'Important', 'Do not use production employees or production attendance for QA. Manual and biometric tests should use a dedicated QA employee and test dates.', fill='FFF4CC')
    doc.core_properties.title = 'Attendance, Shift and Overtime QA SOP'
    doc.core_properties.subject = 'QA procedure for Odoo HR attendance and shift modules'
    doc.core_properties.author = 'QA / HR Operations'
    doc.save(OUT)
    print(OUT.resolve())


if __name__ == '__main__':
    main()
