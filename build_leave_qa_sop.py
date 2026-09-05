from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from build_attendance_qa_sop import (
    BLUE,
    DARK_BLUE,
    INK,
    LIGHT_BLUE,
    LIGHT_GRAY,
    MUTED,
    WHITE,
    WIDTH,
    add_body,
    add_bullet,
    add_callout,
    add_heading,
    add_number,
    add_table,
    add_text,
    configure_styles,
    set_paragraph,
)


OUT = Path('Leave_Management_QA_SOP.docx')


def main():
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header.paragraphs[0]
    set_paragraph(header, after=0, line=1.0, alignment=WD_ALIGN_PARAGRAPH.RIGHT)
    add_text(header, 'HR LEAVE QA  |  CONTROLLED TEST DOCUMENT', size=9, color=MUTED, bold=True)
    footer = section.footer.paragraphs[0]
    set_paragraph(footer, before=0, after=0, line=1.0, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(footer, 'Leave Management QA SOP', size=9, color=MUTED)

    title = doc.add_paragraph()
    set_paragraph(title, before=12, after=3, line=1.0, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(title, 'Leave Management QA SOP', size=24, color=DARK_BLUE, bold=True)
    subtitle = doc.add_paragraph()
    set_paragraph(subtitle, after=16, line=1.0, alignment=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(subtitle, 'Standard operating procedure for functional testing in Odoo', size=11, color=MUTED, italic=True)

    add_table(doc, ['Document control', 'Value'], [
        ('Version', '1.0'),
        ('Owner', 'QA / HR Operations'),
        ('Scope', 'Leave types, policies, allocations, approvals and payroll impact'),
        ('Timezone', 'Asia/Kolkata'),
        ('Test mode', 'Dedicated QA employee and test company'),
    ], [1800, 7560], header_fill=BLUE)

    add_callout(doc, 'Purpose', 'Confirm that leave setup, allocation generation, leave requests, approvals, paid/unpaid behavior, public holidays and employee/payroll results work correctly.')

    add_heading(doc, '1. Preconditions', 1)
    for item in [
        'Use a dedicated QA employee and a test company.',
        'Set the QA user timezone to Asia/Kolkata.',
        'Confirm the employee is active and has an active contract.',
        'Confirm the HR user has leave officer/manager permissions.',
        'Use a draft database or test company so production balances are not affected.',
        'Use a test month with enough working days to verify allocation and payroll behavior.',
    ]:
        add_bullet(doc, item)

    add_heading(doc, '2. Leave type configuration', 1)
    add_table(doc, ['Leave type', 'Code', 'Expected usage'], [
        ('Casual Leave', 'CL', 'Paid leave'),
        ('Sick Time Off', 'SL', 'Paid leave'),
        ('Earned Leaves', 'EL', 'Paid leave'),
        ('Paid Leave', 'PL', 'Paid leave'),
        ('Unpaid', 'UNP', 'Unpaid leave / LOP'),
        ('Maternity Leave', 'MAT', 'Paid or policy-controlled leave'),
        ('Compensatory Days', 'CO', 'Leave earned from approved comp-off'),
    ], [3000, 1200, 5160])
    add_callout(doc, 'Configuration rule', 'Each leave type must have a unique code, the correct attendance usage, the correct allocation requirement, and the correct paid/unpaid setting.')

    add_heading(doc, '3. Allocation policy setup', 1)
    add_body(doc, 'For each allocated leave type, verify the policy and allocation settings before creating requests.')
    add_table(doc, ['Setting', 'QA verification'], [
        ('Leave type', 'Correct leave type is selected'),
        ('Allocation required', 'Enabled where balance control is required'),
        ('Validity', 'Start/end dates are correct'),
        ('Accrual or fixed allocation', 'Matches company policy'),
        ('Employee or department scope', 'Correct scope is selected'),
        ('Carry forward', 'Matches approved policy'),
        ('Approval flow', 'Correct officer/manager approval is required'),
        ('Paid status', 'Paid or unpaid behavior is correct'),
    ], [2700, 6660])

    add_heading(doc, '4. Functional test cases', 1)
    cases = [
        ('LV-001', 'Leave type list and codes',
         ['Open Time Off configuration.', 'Review Casual, Sick, Earned, Paid, Unpaid, Maternity and Compensatory leave types.', 'Open each leave type and record its code and attendance usage.'],
         ['All required leave types exist.', 'Codes are CL, SL, EL, PL, UNP, MAT and CO as configured.', 'No leave type incorrectly uses the GLOBAL code when a specific code is configured.'],
         'Leave type list and form screenshots.'),
        ('LV-002', 'Paid leave configuration',
         ['Open Casual Leave, Sick Time Off and Earned Leaves.', 'Verify each is configured as paid/worked according to policy.', 'Save and reopen the leave types.'],
         ['Paid leave setting is retained.', 'Attendance/payroll usage is correct.', 'The leave type does not create an unpaid deduction.'],
         'Leave form settings.'),
        ('LV-003', 'Unpaid leave configuration',
         ['Open Unpaid leave.', 'Verify code UNP and unpaid/absence behavior.', 'Save and reopen the leave type.'],
         ['Unpaid leave is identified as unpaid.', 'The leave is available for request.', 'The leave can create LOP according to payroll policy.'],
         'Unpaid leave form and code.'),
        ('LV-004', 'Generate allocations',
         ['Open the employee form or leave allocation screen.', 'Click Generate Allocations.', 'Select the required year or policy if available.', 'Review the generated allocation records.'],
         ['Allocations are created for every configured leave type that requires allocation.', 'Allocation dates and balances are correct.', 'No duplicate active allocations are created.'],
         'Before/after allocation list.'),
        ('LV-005', 'Reset allocations',
         ['Create or identify test allocations.', 'Click Reset Allocations.', 'Review the employee allocation list.'],
         ['Only intended generated allocations are reset.', 'Manual or approved historical allocations are not deleted unexpectedly.', 'The result is clearly shown to the user.'],
         'Allocation list before and after reset.'),
        ('LV-006', 'Leave balance visibility',
         ['Open the employee Time Off view.', 'Review each leave balance.', 'Compare the displayed balance with allocation minus approved/requested leave.'],
         ['The employee sees the correct leave types.', 'Available balance is correct.', 'Balance does not become negative unless policy permits it.'],
         'Employee leave balance screenshot.'),
        ('LV-007', 'Paid leave request and approval',
         ['Submit one Casual, Sick or Earned leave request within the allocation period.', 'Approve the request as the authorized approver.', 'Review the employee calendar and attendance result.'],
         ['Request status becomes Approved.', 'Balance is reduced by the approved duration.', 'Leave is shown with the correct code.', 'Paid leave does not create LOP.'],
         'Request, approval and attendance screenshots.'),
        ('LV-008', 'Unpaid leave request and approval',
         ['Submit an Unpaid leave request.', 'Approve the request.', 'Review attendance and draft payroll results.'],
         ['Request status becomes Approved.', 'Balance behavior follows the unpaid policy.', 'Unpaid leave is shown as LOP or unpaid absence.', 'Payable salary is reduced according to the configured formula.'],
         'Request, attendance and payroll screenshots.'),
        ('LV-009', 'Insufficient balance',
         ['Use a leave type with zero or insufficient balance.', 'Request more days than available.', 'Submit the request.'],
         ['The system blocks the request or shows the configured warning.', 'No invalid approval is created.', 'The balance is not corrupted.'],
         'Validation message screenshot.'),
        ('LV-010', 'Overlapping leave dates',
         ['Create an approved leave for a test date.', 'Submit another leave request covering the same date.', 'Try to approve it.'],
         ['Overlapping requests are blocked or handled according to policy.', 'The employee calendar does not double-count the date.', 'Balances remain correct.'],
         'Overlapping request result.'),
        ('LV-011', 'Half-day and multi-day leave',
         ['Submit a half-day leave.', 'Submit a multi-day leave.', 'Approve both requests.'],
         ['Half-day duration is calculated correctly.', 'Multi-day duration excludes configured non-working days where applicable.', 'Attendance and balances match the approved duration.'],
         'Request durations and calendar.'),
        ('LV-012', 'Public holiday and weekly off',
         ['Configure a public holiday and weekly off.', 'Submit leave covering a holiday or weekly off.', 'Review the approved duration.'],
         ['Holiday and weekly-off treatment follows company policy.', 'Leave is not incorrectly counted twice.', 'Paid holiday behavior is correct.'],
         'Calendar and leave duration.'),
        ('LV-013', 'Cancel or refuse leave',
         ['Create an approved leave.', 'Cancel or refuse it according to the workflow.', 'Review the employee balance and calendar.'],
         ['Status changes correctly.', 'Balance is restored when policy requires it.', 'Attendance/payroll values no longer include the cancelled leave.'],
         'Status history and balance.'),
        ('LV-014', 'Payroll leave integration',
         ['Create a draft payslip for a period containing paid leave.', 'Refresh Data and Compute Sheet.', 'Repeat with unpaid leave.'],
         ['Paid leave does not reduce payable basic/HRA for the approved paid duration.', 'Unpaid leave creates the configured LOP deduction.', 'Leave lines have the correct code and duration.'],
         'Worked Days & Inputs and salary computation screenshots.'),
        ('LV-015', 'Permissions and employee access',
         ['Test as employee, HR officer and HR manager.', 'Open leave types, allocations and requests.', 'Attempt unauthorized approval or configuration changes.'],
         ['Employees can submit permitted requests.', 'Officers can review/approve according to policy.', 'Only authorized users can configure policies or generate/reset allocations.'],
         'Role-based access results.'),
        ('LV-016', 'Refresh and regression',
         ['Change an allocation or approve a leave after a draft payslip exists.', 'Click Refresh Data and Compute Sheet.', 'Review the leave lines and balances.'],
         ['Latest approved leave values are loaded.', 'Old values are replaced without deleting the payslip.', 'No duplicate leave lines are created.'],
         'Before/after screenshots and values.'),
    ]
    for case in cases:
        from build_attendance_qa_sop import add_test_case
        add_test_case(doc, *case)

    add_heading(doc, '5. Leave QA result form', 1)
    add_table(doc, ['Field', 'Value'], [
        ('Test Case ID', ''),
        ('Employee', ''),
        ('Leave type / code', ''),
        ('Request dates', ''),
        ('Requested duration', ''),
        ('Allocated balance', ''),
        ('Approved duration', ''),
        ('Expected result', ''),
        ('Actual result', ''),
        ('Status', 'PASS / FAIL'),
        ('Screenshot or evidence', ''),
    ], [3000, 6360])

    add_heading(doc, '6. Defect report format', 1)
    add_table(doc, ['Field', 'Details'], [
        ('Issue title', ''),
        ('Module', 'Leave Type / Policy / Allocation / Request / Payroll'),
        ('Test Case ID', ''),
        ('Employee and leave type', ''),
        ('Steps to reproduce', ''),
        ('Expected result', ''),
        ('Actual result', ''),
        ('Error message', ''),
        ('Screenshot', ''),
        ('Severity', 'Low / Medium / High / Critical'),
    ], [2700, 6660])

    add_callout(doc, 'Important', 'Do not test allocation reset or leave cancellation with production employees. Use a dedicated QA employee and record the original balance before every destructive test.', fill='FFF4CC')
    doc.core_properties.title = 'Leave Management QA SOP'
    doc.core_properties.subject = 'QA procedure for Odoo leave types, allocations and payroll integration'
    doc.core_properties.author = 'QA / HR Operations'
    doc.save(OUT)
    print(OUT.resolve())


if __name__ == '__main__':
    main()
