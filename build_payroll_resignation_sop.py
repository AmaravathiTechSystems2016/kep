# -*- coding: utf-8 -*-

from pathlib import Path

from docx import Document

from build_dotbd_hr_full_sop import (
    BLUE,
    DARK_BLUE,
    GOLD,
    INK,
    LIGHT_BLUE,
    OUTPUT as FULL_OUTPUT,
    add_table,
    bullet,
    callout,
    heading,
    number,
    para,
    set_font,
    style_document,
)


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / 'artifacts' / 'dotbd_hr_payroll_resignation_operator_sop.docx'


def build():
    doc = Document()
    style_document(doc)

    title = doc.add_paragraph(style='Title')
    run = title.add_run('Payroll and Resignation Operations')
    set_font(run, size=27, color=INK, bold=True)
    subtitle = doc.add_paragraph()
    run = subtitle.add_run('Operator SOP for HR, Payroll, Department Managers, and Employees')
    set_font(run, size=15, color=BLUE, bold=True)
    para(doc, 'Version 1.0 | 07 September 2026 | Odoo 19 | Database: kep2')
    callout(doc, 'Purpose', 'Use this guide to configure payroll, process payslips, submit and approve resignations, and complete offboarding without changing historical employee records.')

    heading(doc, '1. Roles and responsibilities', 1)
    add_table(doc, ['Role', 'Payroll responsibility', 'Resignation responsibility'], [
        ['Employee', 'View own payslips and contract-related information.', 'Create and confirm own resignation; view own request.'],
        ['Department Manager', 'Review department payroll information where permitted.', 'Review department employee resignation, handover, and department clearance; approve or reject.'],
        ['HR Officer', 'Configure profiles, contracts, salary rules, and process payslips.', 'Review and approve/reject resignations; complete HR clearance and documents.'],
        ['HR Manager', 'All HR Officer access plus contract and configuration control.', 'Reset approved requests to draft and control final HR processing.'],
        ['Finance', 'Verify deductions, settlement amounts, and payment status.', 'Complete finance clearance and full-and-final settlement inputs.'],
    ], [1900, 3400, 3960])
    callout(doc, 'Access note', 'If Payroll, Contract Overview, or approval buttons are missing, assign the correct HR/Payroll group to the user, then log out and log in again.', GOLD)

    heading(doc, '2. Payroll setup', 1)
    para(doc, 'Payroll has two levels: the shared Payroll Profile stores standard policy values, while the employee contract stores the employee-specific wage and current terms.')
    add_table(doc, ['Record', 'Purpose', 'Where to maintain it'], [
        ['Payroll Profile', 'Shared percentages, allowances, PF/ESI, tax, and overtime policy.', 'Payroll -> Configuration -> Payroll Profiles'],
        ['Employee Contract', 'One employee-specific contract/version with wage and effective dates.', 'Employees -> Employees -> employee -> Payroll'],
        ['Salary Rules', 'Rules that calculate payslip earnings and deductions.', 'Payroll -> Configuration -> Salary Rules'],
        ['Payslip', 'The monthly calculation and final payroll record.', 'Payroll -> Payslips'],
    ], [1900, 3900, 3460])

    heading(doc, '2.1 Maintain Payroll Profiles', 2)
    for step in [
        'Open Payroll -> Configuration -> Payroll Profiles.',
        'Open an existing profile or click New.',
        'Enter the profile name and select its profile code.',
        'Set Basic %, HRA %, DA, Allowance, Washing Allowance, Medical Allowance, Conveyance, Travel, Meal, and Other Allowances.',
        'Set Overtime Multiplier.',
        'Set ESI applicability, employee/employer rates, and wage ceiling.',
        'Set PF rate, PF wage ceiling, and Professional Tax.',
        'Save the profile and use it for new or revised employee contracts.',
    ]:
        number(doc, step)
    callout(doc, 'Important', 'Profile values are database configuration values. Module upgrades do not replace HR-edited profile values. Applying a profile to a contract copies the current profile values into that contract.', GOLD)

    heading(doc, '2.2 Configure an employee contract', 2)
    for step in [
        'Open Employees -> Employees and select the employee.',
        'Open the Payroll tab. HR access is required for the Contract Overview section.',
        'For a new employee, click Load a Template and select the standard contract template.',
        'Set Contract start date, end date if applicable, wage, employee type, contract type, pay category, and working schedule.',
        'Select the Payroll Profile. Confirm the profile values are copied into the contract.',
        'Review statutory details such as UAN, PF/ESI settings, and Professional Tax.',
        'Save the contract.',
    ]:
        number(doc, step)
    add_table(doc, ['Profile', 'Expected standard values'], [
        ['Manufacturing Worker', 'Basic 50%, HRA 30%, Washing ₹1,250, ESI enabled, overtime multiplier 2.'],
        ['Computer Worker With ESI', 'Basic 50%, HRA 30%, Allowance ₹1,250, ESI enabled, overtime multiplier 2.'],
        ['Computer Worker Without ESI', 'Basic 50%, HRA 40%, Medical ₹1,250, ESI disabled, overtime multiplier 2.'],
    ], [3000, 6260])
    para(doc, 'These are the initial profile records. HR can change them from the Payroll Profiles screen. Washing Allowance is normally zero for Computer/Office Staff.')

    heading(doc, '2.3 Salary hike or contract change', 2)
    for step in [
        'Open the employee Payroll tab and click New Contract beside the current contract.',
        'Enter the effective date of the hike and the new wage.',
        'Keep the same Payroll Profile unless the policy itself changes.',
        'Save the new version. Do not overwrite the old contract.',
        'Open More -> History or the History smart button to verify the previous and new versions.',
    ]:
        number(doc, step)
    callout(doc, 'Example', 'If the employee earns ₹30,000 until 30 June and ₹35,000 from 1 July, create a new contract version effective 1 July. Historical payslips continue to use the old version.', LIGHT_BLUE)

    heading(doc, '2.4 Create and process a payslip', 2)
    for step in [
        'Open Payroll -> Payslips and click New.',
        'Select the employee and the correct contract/version.',
        'Set the payslip period and salary structure.',
        'Click Compute Sheet.',
        'Review Basic, HRA, DA, allowances, overtime, PF, ESI, Professional Tax, late check-in deductions, and Net Salary.',
        'Correct contract or input data if a line is wrong, then recompute.',
        'Confirm the payslip and complete payment processing according to company procedure.',
    ]:
        number(doc, step)
    add_table(doc, ['Check', 'Expected result'], [
        ['Basic salary', 'Calculated from wage and the contract Basic Percentage.'],
        ['HRA', 'Calculated from wage and the contract HRA Percentage.'],
        ['Overtime', 'Uses approved overtime input and the contract overtime multiplier.'],
        ['PF', 'Uses PF rate and PF wage ceiling configured on the contract/profile.'],
        ['ESI', 'Deduction applies according to ESI applicability and rate.'],
        ['Late check-in', 'Approved late check-ins for the period appear as a deduction input.'],
        ['Net salary', 'Gross earnings minus applicable deductions.'],
    ], [2600, 6660])

    heading(doc, '3. Resignation and offboarding', 1)
    para(doc, 'A resignation moves through Draft, Confirm, Approved, or Rejected. The employee submits the request, the department manager reviews department work, and HR completes approval and offboarding.')
    add_table(doc, ['Stage', 'Owner', 'Main action'], [
        ['Draft', 'Employee or HR', 'Enter employee, last day, reason, and request details.'],
        ['Confirm', 'Employee', 'Click Confirm to submit the resignation.'],
        ['Review', 'Department Manager', 'Verify handover and department/stores/admin clearance.'],
        ['Approved or Rejected', 'Department Manager / HR', 'Approve or reject after review.'],
        ['Offboarding', 'HR, Finance, Admin', 'Complete settlement, exit interview, documents, and PF/ESI transfer.'],
    ], [2100, 2200, 4960])

    heading(doc, '3.1 Employee submits resignation', 2)
    for step in [
        'Open Resignation -> Resignation Request.',
        'Click New.',
        'Confirm the employee name and department.',
        'Enter the requested Last Day of Employee.',
        'Enter the Reason.',
        'Verify Join Date, Contract Template, and Notice Period.',
        'Save the request and click Confirm.',
    ]:
        number(doc, step)
    callout(doc, 'Date rule', 'The Approved Last Day uses the employee requested Last Day. The Notice Period remains visible for reference and review.', GOLD)

    heading(doc, '3.2 Department Manager review', 2)
    for step in [
        'Log in as the Department Manager and open Resignation.',
        'Open the employee request from the manager’s department.',
        'Review the reason, requested last day, notice period, and employee details.',
        'Set Knowledge Transfer Status to In Progress.',
        'Select Handover To and enter Handover Notes.',
        'After completion, set Knowledge Transfer Status to Completed.',
        'Complete Department, Stores/PPE, and Admin/Asset Clearance.',
        'Click Approve or Reject according to the review result.',
    ]:
        number(doc, step)
    para(doc, 'Department managers should see department employees only. HR users have company-wide resignation access according to their assigned groups.')

    heading(doc, '3.3 HR approval and offboarding', 2)
    for step in [
        'Open the confirmed resignation as an HR Officer or HR Manager.',
        'Select Resignation Type: Normal Resignation or Fired by the company.',
        'Verify the requested last day and approved last day.',
        'Complete HR Clearance, Finance Clearance, Contractor Clearance where applicable, and Clearance Notes.',
        'Click Approve or Reject if not already completed by the department manager.',
        'Complete Full & Final Status, Leave Encashment, Gratuity, Dues Recovery, ESOP Clawback, and Settlement Notes.',
        'Complete Exit Interview and feedback.',
        'Mark Relieving Letter, Experience Certificate, Service Certificate, and Trade Skill Certificate when issued.',
        'Update PF/ESI Transfer status and notes.',
    ]:
        number(doc, step)
    add_table(doc, ['Field group', 'What to enter'], [
        ['Knowledge Transfer', 'Status, handover recipient, and practical handover notes.'],
        ['Exit Clearance', 'Department, Stores/PPE, Admin/Asset, HR, Finance, and Contractor status.'],
        ['Full & Final', 'Status and amounts confirmed by HR/Finance.'],
        ['Exit Interview', 'Completion checkbox and feedback.'],
        ['Relieving Documents', 'Mark each document only after it is issued.'],
        ['PF/ESI Transfer', 'Pending, Assistance Provided, Completed, or Not Applicable, with notes.'],
    ], [2600, 6660])

    heading(doc, '3.4 Employee deactivation', 2)
    para(doc, 'When the approved last day is reached, the scheduled resignation job deactivates the employee and linked user account. If the last day is today or earlier when HR approves the request, deactivation occurs during approval. A daily scheduled action handles future last days.')
    callout(doc, 'Correction procedure', 'If an approved resignation needs correction, an HR Manager can click Set to Draft, update the record, and process it again.', LIGHT_BLUE)

    heading(doc, '4. QA test checklist', 1)
    add_table(doc, ['Test', 'Expected result', 'Pass'], [
        ['Employee creates resignation', 'Employee can create and confirm own request only.', '[ ]'],
        ['Manager visibility', 'Department Manager sees only department employee resignations.', '[ ]'],
        ['Manager approval', 'Manager can review and approve/reject according to assigned HR User access.', '[ ]'],
        ['Requested last day', 'Approved Last Day equals the requested Last Day.', '[ ]'],
        ['Future deactivation', 'Employee becomes inactive when scheduled job reaches the approved date.', '[ ]'],
        ['Payroll profile', 'Profile values are editable from Payroll Profiles.', '[ ]'],
        ['Profile application', 'Changing a profile applies its current UI values to the contract.', '[ ]'],
        ['Salary hike', 'New contract version preserves the old contract and uses the new wage from its effective date.', '[ ]'],
        ['Payslip', 'Computed payslip shows correct earnings, deductions, overtime, and net salary.', '[ ]'],
        ['Employee payroll access', 'Employee reads own payslip only.', '[ ]'],
    ], [3500, 5260, 600])

    heading(doc, '5. Troubleshooting', 1)
    add_table(doc, ['Problem', 'Check / resolution'], [
        ['Contract Template is blank', 'Open the active employee contract and use Load a Template. The contract must have a template assigned.'],
        ['Payroll Profile values do not update', 'Change the profile, save, and verify the contract. Upgrade hr_payroll_community if the UI does not show Payroll Profiles.'],
        ['Previous contract is missing', 'Use New Contract rather than editing the current contract. Open More -> History and verify the version count.'],
        ['Approve button is missing', 'Check Department Manager, HR Officer, or HR Manager group assignment, then log in again.'],
        ['Employee remains active after last day', 'Confirm the resignation is Approved, verify Approved Last Day, and check that the daily scheduled action is running.'],
        ['Payslip amount is wrong', 'Check active contract version, profile values, salary structure, approved overtime, and deduction inputs before recomputing.'],
    ], [2800, 6560])

    heading(doc, '6. Important operating rules', 1)
    for item in [
        'Do not edit a shared Payroll Profile to give one employee a personal raise; create a new employee contract/version.',
        'Do not overwrite an old contract when a salary, role, or contract term changes.',
        'Complete and verify payroll before deactivating a resigned employee when company policy requires final settlement processing.',
        'Use a test employee and test database for QA. Never test approval, deletion, or salary changes on live payroll records without authorization.',
        'Record settlement amounts and document issuance accurately because these fields are used for offboarding tracking, not automatic external payments or document generation.',
    ]:
        bullet(doc, item)
    para(doc, 'End of Payroll and Resignation Operator SOP.')

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == '__main__':
    build()
